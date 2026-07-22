from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import CellIsRule, FormulaRule

# Why does this break the link?
def deleteLinkColumn(sheet, column) -> None:
    #TODO: maybe make this more flexible by having the option to give the column header name to detect position automatically
    if "GeräteId" not in [cell.value for cell in sheet[1]]:
        return
    print("Dropping device Id column...")
    sheet.delete_cols(column)

def formatEmailLinks(sheet) -> None:
    print("Trying to create Email links...")
    link = "https://teams.microsoft.com/l/chat/0/0?users="
    header = [cell.value for cell in sheet[1]]
    emailColumn = 0
    for i, colName in enumerate(header):
        if colName in ["Benutzer-E-Mail"]:
            emailColumn = i

    for row in sheet.rows:
        if row[emailColumn].row == 1:
            continue
        if row[emailColumn].value == None or "NO_EMAIL" in row[emailColumn].value:
            continue
        row[emailColumn].hyperlink = link + row[emailColumn].value
        row[emailColumn].style = "Hyperlink"

def formatOverviewLinks(sheet, keepColumn: bool = False) -> None:
    if "GeräteId" not in [cell.value for cell in sheet[1]]:
        return
    print("Trying to create device overview links...")
    linkList = {}
    for row in sheet.rows:
        if row[0].row == 1:
            continue
        if row[0].value == None or row[0].value == "":
            linkList[row[0].row] = "NO_DEVICE_ID"
        linkList[row[0].row] = f"https://intune.microsoft.com/#view/Microsoft_Intune_Devices/DeviceSettingsMenuBlade/~/properties/mdmDeviceId/{row[0].value}"

    if not keepColumn:
        # If I do this after setting the links, the actual link gets applied to the Email while the device name stays blue...
        deleteLinkColumn(sheet, 1)

    for row in sheet.rows:
        # We assume that the first column is either the DeviceName or the DeviceId if not, we do not care >:3
        if row[0].row == 1:
            continue
        if linkList[row[0].row] == None or "NO_DEVICE_ID" in linkList[row[0].row]:
            continue
        row[0].hyperlink = linkList[row[0].row]
        row[0].style = "Hyperlink"

def columnAutoWidth(sheet) -> None:
    print("Adjusting column widths...")
    # Iterate through each column and adjust the width
    for column in sheet.columns:
        max_length = 0
        max_font = 11
        column_letter = get_column_letter(column[0].column)
        for cell in column:
            try:
                # Measure the length of the cell value
                if len(str(cell.value)) > max_length:
                    max_length = len(cell.value)
                if cell.font.size > max_font:
                    max_font = cell.font.size
            except:
                pass
        # The math accounting the font size is not too important, I just like it more that way
        adjusted_width = (max_length + 2) * (max_font/10)-0.1
        sheet.column_dimensions[column_letter].width = adjusted_width
    pass

def setColorRules(sheet) -> None:
    print("Creating conditional color formating rules...")
    green = PatternFill(start_color="33CC33", end_color="33CC33", fill_type="solid")
    red = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
    blue = PatternFill(start_color="00B0F0", end_color="00B0F0", fill_type="solid")

    header = [cell.value for cell in sheet[1]]

    changedDataColumn = None
    problemColumns = []
    for i, colName in enumerate(header):
        if colName not in ["GeräteId", "Gerätename", "Benutzer-E-Mail", 'Verändert', 'Notes']:
            problemColumns.append(i + 1)
        if colName == 'Verändert':
            changedDataColumn = i + 1

    for col_idx in problemColumns:
        colLetter = get_column_letter(col_idx)
        if sheet.max_row > 1:
            cf_range = f"{colLetter}2:{colLetter}{sheet.max_row}"
            sheet.conditional_formatting.add(cf_range, CellIsRule(operator='equal', formula=['""'], fill=green))
            
            sheet.conditional_formatting.add(cf_range, FormulaRule(formula=[f'_xlfn.REGEXTEST(INDIRECT(ADDRESS(ROW(), COLUMN({colLetter}1))), "^\\s?\\[.*?\\]$")'], fill=green))
            
            sheet.conditional_formatting.add(cf_range, CellIsRule(operator='notEqual', formula=['""'], fill=red))
    
    if changedDataColumn is not None:
        colLetter = get_column_letter(changedDataColumn)
        if sheet.max_row > 1:
            cf_range = f"{colLetter}2:{colLetter}{sheet.max_row}"

            sheet.conditional_formatting.add(cf_range, CellIsRule(operator='equal', formula=['"New"'], fill=green))
            sheet.conditional_formatting.add(cf_range, CellIsRule(operator='equal', formula=['"Removed"'], fill=red))
            sheet.conditional_formatting.add(cf_range, CellIsRule(operator='notEqual', formula=['""'], fill=blue))
            sheet.conditional_formatting.add(cf_range, CellIsRule(operator='equal', formula=['""'], fill=None))
        ...

def addNoteColumn(sheet) -> None:
    #TODO: maybe do this before formating already
    print("Adding 'Notes' column...")
    if "Notes" in [cell.value for cell in sheet[1]]:
        sheet.column_dimensions[get_column_letter(sheet.max_column)].width = 30
        return
    noteColumnNumber = sheet.max_column + 1
    cell = sheet.cell(row=1, column=noteColumnNumber)
    cell.value = "Notes"
    sheet.column_dimensions[get_column_letter(noteColumnNumber)].width = 30

def setHeaderColors(sheet) -> None:
    print("Setting header background colors...")
    gray = PatternFill(start_color="A6A6A6", end_color="A6A6A6", fill_type="solid")
    lightGray = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")

    header = [cell.value for cell in sheet[1]]

    for i, colName in enumerate(header):
        cell = sheet.cell(row=1, column=i+1)
        if colName == "Notes":
            cell.fill = lightGray
        else:
            cell.fill = gray

def convertToTable(sheet, sheetName: str) -> None:
    print("Converting cells into a table...")
    colLetter = get_column_letter(sheet.max_column)
    cellRange = f"A1:{colLetter}{sheet.max_row}"
    tab = Table(displayName=f"Tabelle1_{sheetName}", ref=cellRange)

    style = TableStyleInfo(name="TableStyleLight15", showFirstColumn=True,
                       showLastColumn=False, showRowStripes=True, showColumnStripes=True)
    tab.tableStyleInfo = style

    sheet.add_table(tab)

def formatExcel(filePath: str, sheetName: str = "Windows", keepDeviceIdColumn: bool = True, createDeviceLinks: bool = True, createEmailLinks: bool = False) -> None:
    wb = load_workbook(filePath)
    sheet = wb[sheetName] if sheetName in wb else None

    if not sheet:
        return

    print(f"Formatting excel spreadsheet '{sheetName}' in file '{filePath}'...")

    if createDeviceLinks:
        formatOverviewLinks(sheet, keepColumn=keepDeviceIdColumn)
    elif not createDeviceLinks and not keepDeviceIdColumn:
        deleteLinkColumn(sheet=sheet, column=1)
    if createEmailLinks:
        formatEmailLinks(sheet)

    setColorRules(sheet)
    
    columnAutoWidth(sheet)

    addNoteColumn(sheet)
    setHeaderColors(sheet)
    convertToTable(sheet, sheetName)

    wb.save(filePath)